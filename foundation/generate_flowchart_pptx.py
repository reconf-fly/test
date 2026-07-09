from pptx import Presentation  # 导入PPT对象用于创建演示文稿
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE  # 导入形状枚举用于绘制节点
from pptx.enum.shapes import MSO_CONNECTOR  # 导入连接线枚举用于绘制箭头
from pptx.enum.text import PP_ALIGN  # 导入文本对齐枚举用于居中显示
from pptx.util import Inches  # 导入英寸单位转换工具
from pptx.dml.color import RGBColor  # 导入颜色对象用于设置填充和线条

presentation = Presentation()  # 创建一个新的演示文稿对象
presentation.slide_width = Inches(16)  # 设置页面宽度为16英寸便于横向布局
presentation.slide_height = Inches(9)  # 设置页面高度为9英寸便于容纳流程图
slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # 添加空白页作为流程图画布

title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(15.4), Inches(0.5))  # 创建标题文本框
title_frame = title_box.text_frame  # 获取标题文本框的文本容器
title_frame.clear()  # 清空默认段落避免重复文本
title_paragraph = title_frame.paragraphs[0]  # 获取首段用于写入标题内容
title_paragraph.text = "总流程图：审计知识图谱驱动的规格生成与迭代验证"  # 写入流程图标题文本
title_paragraph.alignment = PP_ALIGN.CENTER  # 设置标题为居中对齐
title_paragraph.runs[0].font.size = Inches(0.22)  # 设置标题字号以保证清晰可读
title_paragraph.runs[0].font.bold = True  # 设置标题为加粗增强视觉层级

nodes = {  # 定义所有节点的位置尺寸和显示文本
    "A": ("历史审计报告与源码", 0.3, 0.9, 2.0, 0.55, "rect"),  # 定义A节点信息
    "B1": ("DeFi语义提取", 0.3, 1.7, 2.0, 0.55, "rect"),  # 定义B1节点信息
    "B2": ("漏洞模式提取", 0.3, 2.5, 2.0, 0.55, "rect"),  # 定义B2节点信息
    "C1": ("去重/抽象/分类", 2.7, 1.7, 2.0, 0.55, "rect"),  # 定义C1节点信息
    "C2": ("去重/抽象/分类", 2.7, 2.5, 2.0, 0.55, "rect"),  # 定义C2节点信息
    "D": ("语义子图", 5.1, 1.7, 1.6, 0.55, "rect"),  # 定义D节点信息
    "E": ("漏洞子图", 5.1, 2.5, 1.6, 0.55, "rect"),  # 定义E节点信息
    "F": ("LLM因果链接", 7.1, 2.1, 1.8, 0.55, "rect"),  # 定义F节点信息
    "G": ("审计知识图谱G", 9.3, 1.8, 2.0, 0.55, "rect"),  # 定义G节点信息
    "Gs": ("链接置信度评分", 9.3, 2.6, 2.0, 0.55, "rect"),  # 定义Gs节点信息
    "N": ("新Solidity项目", 0.3, 4.0, 2.0, 0.55, "rect"),  # 定义N节点信息
    "M": ("Knowledge Mapper", 5.1, 4.0, 2.3, 0.55, "rect"),  # 定义M节点信息
    "R": ("归一化加权排序器\n(借鉴PropertyGPT多维评分)", 7.8, 4.0, 3.2, 0.85, "rect"),  # 定义R节点信息
    "K": ("Top-K语义-漏洞对", 11.4, 4.1, 2.2, 0.55, "rect"),  # 定义K节点信息
    "S": ("规格生成器 Spec Generator", 11.4, 5.0, 2.2, 0.55, "rect"),  # 定义S节点信息
    "P": ("候选规格重排\n(可编译性+相似度评分)", 11.4, 5.9, 2.2, 0.85, "rect"),  # 定义P节点信息
    "H": ("Harness合成器", 11.4, 7.0, 2.2, 0.55, "rect"),  # 定义H节点信息
    "X": ("Fuzz/执行器\n(Foundry + Coverage + Sanitizer)", 8.5, 7.0, 2.5, 0.85, "rect"),  # 定义X节点信息
    "Y": ("反思器 Reflector\n有效性判定+失败归因", 5.6, 7.0, 2.5, 0.85, "rect"),  # 定义Y节点信息
    "W": ("工作记忆 Working Memory\n覆盖率/trace/状态变化/失败原因", 2.4, 7.0, 2.8, 0.9, "rect"),  # 定义W节点信息
    "Z": ("触发迭代?", 2.7, 5.9, 1.8, 0.85, "diamond"),  # 定义Z判定节点信息
    "Knext": ("下一个语义-漏洞对", 0.3, 5.9, 2.1, 0.55, "rect"),  # 定义Knext节点信息
    "O1": ("Propose: 重写Harness DSL\nA/G/Sigma/Phi/Psi", 0.3, 4.9, 2.3, 0.85, "rect"),  # 定义O1节点信息
    "O2": ("静态验证\n格式/变量绑定/图连通", 0.3, 4.0, 2.3, 0.85, "rect"),  # 定义O2节点信息
    "O3": ("Execute & Observe", 2.9, 4.9, 2.0, 0.55, "rect"),  # 定义O3节点信息
    "O4": ("Score", 2.9, 4.2, 2.0, 0.55, "rect"),  # 定义O4节点信息
    "O5": ("Diagnose", 2.9, 3.5, 2.0, 0.55, "rect"),  # 定义O5节点信息
    "O6": ("更新历史档案\n避免重复试错", 2.9, 2.6, 2.0, 0.85, "rect"),  # 定义O6节点信息
    "Q": ("是否真实漏洞", 5.6, 5.6, 2.0, 0.85, "diamond"),  # 定义Q判定节点信息
    "V": ("报告漏洞+PoC", 5.6, 4.6, 2.0, 0.55, "rect"),  # 定义V节点信息
    "U": ("回灌知识图谱G", 7.9, 4.6, 2.0, 0.55, "rect"),  # 定义U节点信息
    "T": ("记录并丢弃/保留为反例", 5.6, 6.7, 2.0, 0.55, "rect"),  # 定义T节点信息
}  # 结束节点定义

shape_map = {}  # 创建字典用于保存节点ID和PPT形状对象映射

for node_id, (text, left, top, width, height, kind) in nodes.items():  # 遍历所有节点并创建图形
    shape_type = MSO_AUTO_SHAPE_TYPE.DIAMOND if kind == "diamond" else MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE  # 根据节点类型选择形状
    shape = slide.shapes.add_shape(shape_type, Inches(left), Inches(top), Inches(width), Inches(height))  # 在指定位置添加形状
    fill = shape.fill  # 获取形状填充对象
    fill.solid()  # 将填充设置为纯色填充
    fill.fore_color.rgb = RGBColor(235, 244, 255) if kind == "rect" else RGBColor(255, 242, 204)  # 根据类型设置不同底色
    shape.line.color.rgb = RGBColor(43, 87, 154)  # 设置边框颜色为蓝色
    shape.line.width = Inches(0.01)  # 设置边框宽度增强边界可见性
    text_frame = shape.text_frame  # 获取形状文本框对象
    text_frame.clear()  # 清除默认占位文本
    paragraph = text_frame.paragraphs[0]  # 获取首段以写入节点文本
    paragraph.text = text  # 写入节点文字内容
    paragraph.alignment = PP_ALIGN.CENTER  # 将节点文本设置为居中
    paragraph.runs[0].font.size = Inches(0.12)  # 设置节点字体大小保证密集布局可读
    shape_map[node_id] = shape  # 保存节点ID与形状对象关系

edges = [  # 定义节点之间的连接关系以及可选标签
    ("A", "B1", ""), ("A", "B2", ""), ("B1", "C1", ""), ("B2", "C2", ""),  # 定义上游提取与分类链路
    ("C1", "D", ""), ("C2", "E", ""), ("D", "F", ""), ("E", "F", ""),  # 定义子图聚合到因果链接
    ("F", "G", ""), ("F", "Gs", ""), ("N", "M", ""), ("G", "M", ""), ("Gs", "M", ""),  # 定义知识图谱注入映射器
    ("M", "R", ""), ("R", "K", ""), ("K", "S", ""), ("S", "P", ""), ("P", "H", ""),  # 定义候选规格主链路
    ("H", "X", ""), ("X", "Y", ""), ("Y", "W", ""), ("W", "Z", ""),  # 定义执行反馈到判定节点
    ("Z", "Knext", "否"), ("Knext", "S", ""), ("Z", "O1", "是"), ("O1", "O2", ""),  # 定义未触发迭代和触发迭代路径
    ("O2", "O3", ""), ("O3", "O4", ""), ("O4", "O5", ""), ("O5", "O6", ""), ("O6", "H", ""),  # 定义迭代内循环路径
    ("Y", "Q", ""), ("Q", "V", "是"), ("V", "U", ""), ("Q", "T", "否"),  # 定义漏洞真实性判定分支
]  # 结束边定义

def shape_center(shape):  # 定义函数用于计算形状中心点坐标
    return shape.left + shape.width // 2, shape.top + shape.height // 2  # 返回中心点的横纵坐标

for start_id, end_id, label in edges:  # 遍历所有边并绘制连接线
    start_shape = shape_map[start_id]  # 获取起点节点形状对象
    end_shape = shape_map[end_id]  # 获取终点节点形状对象
    start_x, start_y = shape_center(start_shape)  # 计算起点中心坐标
    end_x, end_y = shape_center(end_shape)  # 计算终点中心坐标
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y)  # 绘制直线连接器
    connector.line.color.rgb = RGBColor(79, 79, 79)  # 设置连接线颜色为深灰色
    connector.line.width = Inches(0.008)  # 设置连接线宽度
    connector.line.end_arrowhead = True  # 设置连接线末端箭头以表示方向
    if label:  # 判断当前连接是否需要显示标签文本
        label_left = (start_x + end_x) // 2  # 计算标签横向中点
        label_top = (start_y + end_y) // 2  # 计算标签纵向中点
        label_box = slide.shapes.add_textbox(label_left, label_top, Inches(0.35), Inches(0.2))  # 添加标签文本框
        label_frame = label_box.text_frame  # 获取标签文本框容器
        label_frame.clear()  # 清空标签默认段落
        label_paragraph = label_frame.paragraphs[0]  # 获取标签首段
        label_paragraph.text = label  # 设置标签文本内容
        label_paragraph.alignment = PP_ALIGN.CENTER  # 设置标签文字居中
        label_paragraph.runs[0].font.size = Inches(0.09)  # 设置标签字体大小
        label_paragraph.runs[0].font.bold = True  # 设置标签为加粗突出分支语义

output_path = "overall-flowchart.pptx"  # 定义输出PPT文件名
presentation.save(output_path)  # 保存演示文稿到仓库根目录
print(f"已生成: {output_path}")  # 打印生成结果便于终端确认
